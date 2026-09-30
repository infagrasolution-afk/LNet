import os
import io
import re
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TEMPLATE_DIR = os.path.join(os.path.dirname(__file__), "templates")
INDIVIDUAL_TEMPLATE_PATH = os.path.join(TEMPLATE_DIR, "presupuesto_individual_template.xlsx")

def _extract_numeric(val: Any, default: float = 0.0) -> float:
    """Extrae el primer número de una cadena (ej. '150 mts', '2 unid', '1') de forma segura."""
    if val is None:
        return default
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", ".")
    match = re.search(r"[-+]?\d*\.?\d+", s)
    if match:
        try:
            return float(match.group())
        except ValueError:
            return default
    return default

def calculate_fiber_subtotal(meters: float) -> float:
    """
    Calcula el monto del tendido de fibra según la regla oficial de NetUno:
    - Hasta 120 metros: Tarifa plana de $40.00
    - Metros excedentes (> 120m): $0.40 por cada metro adicional
    """
    if meters <= 0:
        return 0.0
    if meters <= 120:
        return 40.0
    return 40.0 + (meters - 120) * 0.40

def calculate_record_subtotal(record: Dict[str, Any]) -> float:
    """Calcula el total base en USD para una planilla / ticket."""
    activities = record.get("activities", [])
    subtotal = 0.0
    has_fiber = False

    for act in activities:
        if not act.get("checked"):
            continue
        name = (act.get("name") or "").lower()
        desc = (act.get("description") or "").lower()
        qty = _extract_numeric(act.get("unid_mts"), default=1.0)

        if "fibra" in name or "fibra" in desc or act.get("id") == "act-1":
            subtotal += calculate_fiber_subtotal(qty)
            has_fiber = True
        elif "tuberia metalica" in name or "tuberia metalica" in desc:
            subtotal += qty * 10.0
        elif "tuberia plastica" in name or "tuberia plastica" in desc:
            subtotal += qty * 15.0

    # Si no se marcó fibra pero se cargó la solicitud, fijar base mínima estándar de $40 si tiene actividades
    if not has_fiber and any(act.get("checked") for act in activities):
        subtotal += 40.0

    return round(subtotal if subtotal > 0 else 40.0, 2)

# ==========================================
# 1. GENERADOR DE PRESUPUESTO INDIVIDUAL (NETUNO)
# ==========================================

def generate_netuno_individual_excel(record: Dict[str, Any]) -> io.BytesIO:
    """
    Genera el formato individual de presupuesto/factura para un ticket específico
    a partir de la plantilla corporativa de Contratistas NetUno.
    """
    if os.path.exists(INDIVIDUAL_TEMPLATE_PATH):
        wb = openpyxl.load_workbook(INDIVIDUAL_TEMPLATE_PATH)
        ws = wb.active
    else:
        # Fallback si no existe la plantilla
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Presupuesto"

    solicitud = str(record.get("solicitud_num", "")).strip()
    client_name = record.get("client_name", "").strip()
    created_at = record.get("created_at", "")
    date_part = created_at[:10] if len(created_at) >= 10 else datetime.now().strftime("%Y-%m-%d")
    date_num = date_part.replace("-", "")

    # Título de pestaña
    sheet_title = f"P_{solicitud}"[:31]
    ws.title = sheet_title

    # Cabecera principal
    ws["E22"] = client_name
    ws["E23"] = solicitud
    ws["M18"] = f"{date_num}-S{solicitud}"
    ws["M19"] = date_part

    # Representante / Técnico
    tech_name = record.get("created_by_name") or record.get("created_by") or "Angelo J Requena T"
    ws["M20"] = tech_name

    # Actividades y Materiales
    activities = record.get("activities", [])
    fiber_qty = 0.0
    material_counts = {
        "conector": 0.0,
        "roseta": 0.0,
        "hibrido": 0.0,
        "ont": 0.0,
        "tirrap": 0.0,
        "tensor": 0.0,
        "gancho": 0.0,
    }

    extra_items = []

    for act in activities:
        if not act.get("checked"):
            continue
        name = (act.get("name") or "").lower()
        desc = (act.get("description") or "").lower()
        qty = _extract_numeric(act.get("unid_mts"), default=1.0)

        if "fibra" in name or "fibra" in desc or act.get("id") == "act-1":
            fiber_qty += qty
        elif "conector" in name or "conector" in desc or act.get("id") == "act-3":
            material_counts["conector"] += qty
        elif "roseta" in name or "roseta" in desc or act.get("id") == "act-2":
            material_counts["roseta"] += qty
        elif "hibrido" in name or "patch cord" in name or "patch cord" in desc or act.get("id") == "act-4":
            material_counts["hibrido"] += qty
        elif "ont" in name or "ont" in desc or act.get("id") == "act-9":
            material_counts["ont"] += qty
        elif "etiqueta" in name or "etiqueta" in desc or "tirrap" in name or act.get("id") == "act-8":
            material_counts["tirrap"] += qty
        elif "tensor" in name or "tensor" in desc or act.get("id") == "act-5":
            material_counts["tensor"] += qty
        elif "gancho" in name or "gancho" in desc or act.get("id") == "act-6":
            material_counts["gancho"] += qty
        else:
            # Otros materiales o tuberías
            extra_items.append((act.get("description") or act.get("name"), qty))

    # Cargar tendido de fibra en D25
    ws["D25"] = fiber_qty
    ws["L25"] = 0.40
    ws["M25"] = "=IF(D25<=120,40,40+(L25*(D25-120)))"

    # Fila 28: Materiales consumidos en cabecera
    ws["E28"] = material_counts["conector"]
    ws["F28"] = material_counts["roseta"]
    ws["G28"] = material_counts["hibrido"]
    ws["H28"] = material_counts["ont"]
    ws["I28"] = material_counts["tirrap"]
    ws["J28"] = material_counts["tensor"]
    ws["K28"] = material_counts["gancho"]

    # Fórmulas de Subtotal, IVA y Total Factura
    ws["M49"] = "=SUM(M25:M48)"
    ws["M51"] = "=SUM(M49*16%)"
    ws["M53"] = "=SUM(M49:M52)"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# ==========================================
# 2. GENERADOR DE RELACIÓN CONSOLIDADA (NETUNO)
# ==========================================

def generate_netuno_relacion_excel(
    records: List[Dict[str, Any]],
    start_date: Optional[str] = None,
    end_date: Optional[str] = None
) -> io.BytesIO:
    """
    Genera el reporte consolidado 'Relación de Presupuestos' matching exacto
    con el formato contable oficial de contratistas NetUno.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CORP_INST"
    ws.views.sheetView[0].showGridLines = True

    # Estilos oficiales
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=9, bold=False, color="000000")
    total_font = Font(name="Arial", size=10, bold=True, color="000000")

    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC")
    )
    total_border = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000")
    )

    # Título superior / Banner
    period_text = f"Período: {start_date or 'Inicio'} al {end_date or 'Actualidad'}"
    ws.merge_cells("B1:N1")
    ws["B1"] = f"RELACIÓN DE PRESUPUESTOS Y SERVICIOS EJECUTADOS — {period_text.upper()}"
    ws["B1"].font = Font(name="Arial", size=12, bold=True, color="1F497D")
    ws["B1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Encabezados en Fila 3
    headers = [
        ("B", "CONTRATISTAS", 20),
        ("C", "NÚMERO DE PRESUPUESTOS", 24),
        ("D", "TIPO DE INSTALACIÓN", 20),
        ("E", "FECHA", 14),
        ("F", "TICKET", 14),
        ("G", "DESCRIPCIÓN (CLIENTE)", 32),
        ("H", "TOTAL $", 14),
        ("I", "IVA $", 14),
        ("J", "TOTAL FACTURA $", 16),
        ("K", "RET ISLR $", 14),
        ("L", "RET IMPUESTO MUNICIPAL", 16),
        ("M", "RET IVA $", 14),
        ("N", "TOTAL NETO A PAGAR $", 18),
        ("O", "OBSERVACIÓN", 30),
        ("P", "PAGO REALIZADO", 16),
        ("Q", "FECHA DEL PAGO", 16),
    ]

    ws.row_dimensions[3].height = 28
    for col_letter, title, width in headers:
        cell = ws[f"{col_letter}3"]
        cell.value = title
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[col_letter].width = width

    start_row = 4
    current_row = start_row

    for rec in records:
        solicitud = str(rec.get("solicitud_num", "")).strip()
        client_name = rec.get("client_name", "").strip()
        created_at = rec.get("created_at", "")
        date_part = created_at[:10] if len(created_at) >= 10 else datetime.now().strftime("%Y-%m-%d")
        date_num = date_part.replace("-", "")
        budget_num = f"{date_num}-S{solicitud}"

        # Subtotal de la instalación
        subtotal = calculate_record_subtotal(rec)

        # Fila de datos
        ws[f"B{current_row}"] = "INSTALACIONES FP"
        ws[f"C{current_row}"] = budget_num
        ws[f"D{current_row}"] = "EDIFICIO / CORP"
        ws[f"E{current_row}"] = date_part
        ws[f"F{current_row}"] = solicitud
        ws[f"G{current_row}"] = client_name
        ws[f"H{current_row}"] = subtotal
        ws[f"I{current_row}"] = f"=H{current_row}*16%"
        ws[f"J{current_row}"] = f"=H{current_row}+I{current_row}"
        ws[f"K{current_row}"] = f"=H{current_row}*2%"
        ws[f"L{current_row}"] = 0.0
        ws[f"M{current_row}"] = f"=I{current_row}*75%"
        ws[f"N{current_row}"] = f"=J{current_row}-K{current_row}-M{current_row}"
        ws[f"O{current_row}"] = rec.get("observations") or ""
        ws[f"P{current_row}"] = ""
        ws[f"Q{current_row}"] = ""

        # Aplicar formatos y bordes
        for col_letter, _, _ in headers:
            c = ws[f"{col_letter}{current_row}"]
            c.font = data_font
            c.border = thin_border
            if col_letter in ["B", "C", "D", "E", "F", "P", "Q"]:
                c.alignment = center_align
            elif col_letter in ["G", "O"]:
                c.alignment = left_align
            elif col_letter in ["H", "I", "J", "K", "L", "M", "N"]:
                c.alignment = right_align
                c.number_format = "$#,##0.00"

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # Fila de Totales si hay registros
    if current_row > start_row:
        last_data_row = current_row - 1
        ws[f"G{current_row}"] = "TOTALES GENERALES:"
        ws[f"G{current_row}"].font = total_font
        ws[f"G{current_row}"].alignment = right_align

        for col in ["H", "I", "J", "K", "L", "M", "N"]:
            c = ws[f"{col}{current_row}"]
            c.value = f"=SUM({col}{start_row}:{col}{last_data_row})"
            c.font = total_font
            c.alignment = right_align
            c.number_format = "$#,##0.00"
            c.border = total_border

        ws.row_dimensions[current_row].height = 24

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
